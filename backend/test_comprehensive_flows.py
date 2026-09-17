import os
import sys
from io import BytesIO
from pathlib import Path

# Ensure backend root and project root are in sys.path
backend_dir = os.path.dirname(os.path.abspath(__file__))
project_dir = os.path.dirname(backend_dir)
sys.path.insert(0, backend_dir)
sys.path.insert(0, project_dir)

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import fitz  # PyMuPDF
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from main import app
from app.services.ocr_service import configure_tesseract, extract_text_from_image, TESSERACT_INSTALL_GUIDE
from app.services.chat_service import (
    generate_answer,
    load_document,
    load_documents,
    require_client,
    UNAVAILABLE_MESSAGE,
    find_text_span,
)
from app.services.retrieval_service import retrieve_relevant_chunks

client = TestClient(app)

print('=' * 60)
print('DOCUMIND AI - COMPREHENSIVE FLOW VERIFICATION SUITE')
print('=' * 60)

test_results = {}

# Clean up any leftover test docs from prior runs
try:
    existing_docs = client.get('/api/documents').json()
    for d in existing_docs:
        doc_id = d.get('document_id') or d.get('id')
        if d.get('filename') in ['sample_test_documind.pdf', 'sample_test_documind.docx', 'sample_invoice_receipt.png'] and doc_id:
            client.delete(f"/api/documents/{doc_id}")
except Exception as e:
    print(f"Pre-cleanup notice: {e}")

# 1. OCR Configuration Verification
print('\n[STEP 1] Testing Tesseract OCR Configuration...')
tesseract_ready = configure_tesseract()
print(f' -> configure_tesseract(): {tesseract_ready}')
assert tesseract_ready, 'Tesseract must be configured and available'
test_results['Tesseract OCR Configuration'] = 'PASS'

# 2. Image OCR Direct Text Extraction
print('\n[STEP 2] Testing Direct Image OCR Extraction...')
img = Image.new('RGB', (800, 200), color='white')
draw = ImageDraw.Draw(img)
test_image_text = 'DocuMind Invoice ID: INV-2026-999. Total Price: $850.00.'
draw.text((20, 80), test_image_text, fill='black')

buf = BytesIO()
img.save(buf, format='PNG')
png_bytes = buf.getvalue()

extracted_ocr = extract_text_from_image(png_bytes)
print(f' -> Extracted OCR text: {extracted_ocr.strip()}')
assert 'INV-2026-999' in extracted_ocr or '850' in extracted_ocr, f'Image OCR missing expected text: {extracted_ocr}'
test_results['Image OCR Extraction'] = 'PASS'

# 3. PDF Upload Flow
print('\n[STEP 3] Testing PDF Document Upload...')
doc_pdf = fitz.open()
page = doc_pdf.new_page()
pdf_text_content = (
    'Artificial Intelligence and Deep Learning in DocuMind.\n'
    'DocuMind provides neural document retrieval across enterprise datasets.\n'
    'The primary retrieval algorithm uses dense semantic embeddings and lexical BM25.'
)
page.insert_text((50, 100), pdf_text_content, fontsize=12)
pdf_bytes = doc_pdf.tobytes()
doc_pdf.close()

pdf_filename = 'sample_test_documind.pdf'
pdf_res = client.post(
    '/api/upload',
    files={'files': (pdf_filename, pdf_bytes, 'application/pdf')},
)
print(f' -> PDF Upload Status: {pdf_res.status_code}')
assert pdf_res.status_code == 200, f'PDF upload failed: {pdf_res.text}'
pdf_data = pdf_res.json()
pdf_doc_id = pdf_data['document_id']
assert pdf_data['total_pages'] == 1
assert pdf_data['total_chunks'] >= 1
print(f' -> PDF Uploaded Document ID: {pdf_doc_id}')
test_results['PDF Upload Flow'] = 'PASS'

# 4. Duplicate Upload / HTTP 409 Handling
print('\n[STEP 4] Testing Duplicate Upload (HTTP 409)...')
dup_res = client.post(
    '/api/upload',
    files={'files': (pdf_filename, pdf_bytes, 'application/pdf')},
)
print(f' -> Duplicate Upload Status: {dup_res.status_code}')
assert dup_res.status_code == 409, f'Expected 409, got: {dup_res.status_code}'
print(' -> Duplicate upload properly rejected with 409!')
test_results['Duplicate Upload (HTTP 409)'] = 'PASS'

# 5. DOCX Upload Flow
print("\n[STEP 5] Testing DOCX Document Upload...")
docx = Document()
docx.add_heading("DocuMind System Specifications", level=1)
docx.add_paragraph("DocuMind AI processes Word documents by parsing paragraph-level segments.")
docx.add_paragraph("The memory storage layer caches conversation histories in JSON format.")
docx_buf = BytesIO()
docx.save(docx_buf)
docx_bytes = docx_buf.getvalue()

docx_filename = "sample_test_documind.docx"
docx_res = client.post(
    "/api/upload",
    files={"files": (docx_filename, docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
)
print(f" -> DOCX Upload Status: {docx_res.status_code}")
assert docx_res.status_code == 200, f'DOCX upload failed: {docx_res.text}'
docx_data = docx_res.json()
docx_doc_id = docx_data['document_id']
assert docx_data['total_chunks'] >= 1
print(f' -> DOCX Uploaded Document ID: {docx_doc_id}')
test_results['DOCX Upload Flow'] = 'PASS'

# 6. PNG Image Upload Flow (with OCR)
print('\n[STEP 6] Testing PNG Image Upload with OCR.')
image_filename = 'sample_invoice_receipt.png'
img_res = client.post(
    '/api/upload',
    files={'files': (image_filename, png_bytes, 'image/png')},
)
print(f' -> Image Upload Status: {img_res.status_code}')
assert img_res.status_code == 200, f'Image upload failed: {img_res.text}'
img_data = img_res.json()
img_doc_id = img_data['document_id']
assert img_data['file_type'] == 'png'
assert img_data['total_pages'] == 1
assert img_data['total_chunks'] >= 1
print(f' -> Image Uploaded Document ID: {img_doc_id}')
test_results['Image Upload Flow (OCR)'] = 'PASS'

# 7. Document File Download API
print('\n[STEP 7] Testing Document File Download API...')
file_res = client.get(f'/api/documents/{img_doc_id}/file')
assert file_res.status_code == 200
assert file_res.headers.get('content-type') == 'image/png'
assert len(file_res.content) == len(png_bytes)
print(' -> Image File Download API: PASS')

doc_details = client.get(f'/api/documents/{img_doc_id}')
assert doc_details.status_code == 200
assert doc_details.json()['filename'] == image_filename
print(' -> Document Details API: PASS')
test_results['Document File Download API'] = 'PASS'

# 8. Question About Image Text
print('\n[STEP 8] Testing Question About Image Text...')
img_doc = load_document(img_doc_id)
retrieved = retrieve_relevant_chunks(
    document=img_doc,
    question='What is the invoice ID and total price?',
    document_id=img_doc_id,
    max_results=2,
)
print(f' -> Retrieved {len(retrieved)} chunk(s) from image document')
assert len(retrieved) > 0, 'Expected image chunk to be retrieved'
assert 'INV-2026-999' in retrieved[0]['text'] or '850' in retrieved[0]['text']

# Test live API call to xAI
from unittest.mock import patch, MagicMock
from app.services.chat_service import llm_client

try:
    live_res = generate_answer(
        question='What is the invoice ID?',
        document=img_doc,
        retrieved_chunks=retrieved,
        document_ids=[img_doc_id],
        documents=[img_doc],
    )
    print(f' -> Live API Answer: {live_res.get("answer")[:50]}...')
except Exception as err:
    print(f' -> Live API call exception note: {err}')

# Test full pipeline with mocked completion
import json
mock_resp = MagicMock()
mock_choice = MagicMock()
mock_choice.message.content = json.dumps({"answer": "The invoice ID is INV-2026-999 and the total price is $850.00."})
mock_resp.choices = [mock_choice]

with patch.object(llm_client.chat.completions, 'create', return_value=mock_resp):
    mock_answer_result = generate_answer(
        question='What is the invoice ID?',
        document=img_doc,
        retrieved_chunks=retrieved,
        document_ids=[img_doc_id],
        documents=[img_doc],
    )
    print(f' -> Image Q&A Sources Count: {len(mock_answer_result.get("sources", []))}')
    assert len(mock_answer_result.get('sources', [])) > 0, 'Expected image sources'
    img_source = mock_answer_result['sources'][0]
    print(f' -> Source title: {img_source["title"]}')
    print(f' -> Source preview: {img_source["preview"]}')
    assert 'Image' in img_source['title'] or 'sample_invoice_receipt' in img_source['title']
test_results['Question About Image Text'] = 'PASS'

# 9. Question About PDF Text
print('\n[STEP 9] Testing Question About PDF Text...')
pdf_doc = load_document(pdf_doc_id)
pdf_retrieved = retrieve_relevant_chunks(
    document=pdf_doc,
    question='What does DocuMind AI rely on for retrieval?',
    document_id=pdf_doc_id,
    max_results=2,
)
print(f' -> Retrieved {len(pdf_retrieved)} chunk(s) from PDF document')
assert len(pdf_retrieved) > 0
assert 'semantic embeddings' in pdf_retrieved[0]['text'] or 'retrieval' in pdf_retrieved[0]['text'].lower()

mock_pdf_choice = MagicMock()
mock_pdf_choice.message.content = json.dumps({"answer": "DocuMind AI uses dense semantic embeddings and lexical BM25 for hybrid retrieval."})
mock_pdf_resp = MagicMock()
mock_pdf_resp.choices = [mock_pdf_choice]

with patch.object(llm_client.chat.completions, 'create', return_value=mock_pdf_resp):
    pdf_answer_result = generate_answer(
        question='What does DocuMind AI rely on for retrieval?',
        document=pdf_doc,
        retrieved_chunks=pdf_retrieved,
        document_ids=[pdf_doc_id],
        documents=[pdf_doc],
    )
    assert len(pdf_answer_result.get('sources', [])) > 0
    print(f' -> PDF Source: {pdf_answer_result["sources"][0]["title"]}')
    assert 'Page 1' in pdf_answer_result["sources"][0]['title']
test_results['Question About PDF Text'] = 'PASS'

# 10. Multi-Document Question (PDF + Image)
print('\n[STEP 10] Testing Multi-Document Question (PDF + Image)...')
multi_docs = load_documents([pdf_doc_id, img_doc_id])
assert len(multi_docs) == 2

multi_retrieved = retrieve_relevant_chunks(
    document=multi_docs,
    question='What is the invoice price and how does neural document retrieval work?',
    document_id=f'{pdf_doc_id}|{img_doc_id}',
    max_results=4,
)
print(f' -> Multi-doc retrieved {len(multi_retrieved)} chunk(s)')
assert len(multi_retrieved) > 0

mock_multi_choice = MagicMock()
mock_multi_choice.message.content = json.dumps({"answer": "The invoice price is $850.00 and neural document retrieval uses semantic embeddings."})
mock_multi_resp = MagicMock()
mock_multi_resp.choices = [mock_multi_choice]

with patch.object(llm_client.chat.completions, 'create', return_value=mock_multi_resp):
    multi_answer = generate_answer(
        question='What is the invoice price and how does neural document retrieval work?',
        document=multi_docs,
        retrieved_chunks=multi_retrieved,
        document_ids=[pdf_doc_id, img_doc_id],
        documents=multi_docs,
    )
    print(f' -> Multi-doc sources count: {len(multi_answer.get("sources", []))}')
    assert len(multi_answer.get('sources', [])) > 0
test_results['Multi-Document Question'] = 'PASS'

# 11. PDF Source Navigation and Highlighting
print('\n[STEP 11] Testing PDF Source Highlighting Logic...')
from frontend.app import render_pdf_page
page_img, highlighted = render_pdf_page(
    pdf_bytes=pdf_bytes,
    page_number=1,
    highlight_text='DocuMind provides neural document retrieval across enterprise datasets. The primary retrieval algorithm uses dense semantic embeddings and lexical BM25.',
)
assert page_img is not None, 'Page image should be rendered'
print(f' -> Highlighting applied: {highlighted}')
assert highlighted is True, 'Expected highlight to be applied to PDF text'
test_results['PDF Source Highlighting'] = 'PASS'

# 12. Missing Information Handling
print('\n[STEP 12] Testing Missing Information Handling...')
no_chunks_result = generate_answer(
    question='What is the recipe for chocolate cake?',
    document=pdf_doc,
    retrieved_chunks=[],
    document_ids=[pdf_doc_id],
    documents=[pdf_doc],
)
print(f' -> Empty retrieval answer: {no_chunks_result["result"]}' if 'result' in no_chunks_result else f' -> Empty retrieval answer: {no_chunks_result["answer"]}')
assert no_chunks_result['answer'] == UNAVAILABLE_MESSAGE
test_results['Missing Information Handling'] = 'PASS'

# 13. Missing XAI_API_KEY Handling
print('\n[STEP 13] Testing Missing XAI_API_KEY Handling...')
saved_xai = os.environ.get('XAI_API_KEY')
saved_groq = os.environ.get('GROQ_API_KEY')
try:
    os.environ['XAI_API_KEY'] = ''
    os.environ['GROQ_API_KEY'] = ''
    error_caught = False
    try:
        require_client()
    except RuntimeError as re_err:
        error_caught = True
        print(f' -> Clean RuntimeError caught: {re_err}')
        assert 'XAI_API_KEY is missing' in str(re_err)
    assert error_caught, 'Expected require_client to raise RuntimeError when API key is missing'
finally:
    if saved_xai is not None:
        os.environ['XAI_API_KEY'] = saved_xai
    if saved_groq is not None:
        os.environ['GROQ_API_KEY'] = saved_groq

test_results['Missing XAI_API_KEY Handling'] = 'PASS'

# 14. Unsupported File Type Error Handling
print('\n[STEP 14] Testing Unsupported File Type Handling...')
bad_res = client.post(
    '/api/upload',
    files={'files': ('test_bad_file.txt', b'Hello world text file', 'text/plain')},
)
print(f' -> Bad file status: {bad_res.status_code}')
assert bad_res.status_code == 400
assert 'Unsupported file type' in bad_res.json().get('detail', '')
test_results['Unsupported File Type Handling'] = 'PASS'

# Cleanup test documents
print('\n[CLEANUP] Removing test artifacts from workspace...')
for doc_id in [pdf_doc_id, docx_doc_id, img_doc_id]:
    del_res = client.delete(f'/api/documents/{doc_id}')
    assert del_res.status_code == 200

print('\n' + '=' * 60)
print('ALL VERIFICATION SUITE TESTS PASSED!')
print('=' * 60)
for test_name, status in test_results.items():
    print(f'  [✔] {test_name}: {status}')
print('=' * 60)
