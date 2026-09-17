import os
import sys
from io import BytesIO
from pathlib import Path

# Set stdout encoding for Windows console
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

backend_dir = Path(__file__).resolve().parent
sys.path.insert(0, str(backend_dir))

import fitz  # PyMuPDF
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from main import app

client = TestClient(app)

print("=" * 65)
print("DOCUMIND AI - CHAT FLOW & OCR TEST SUITE")
print("=" * 65)

test_results = {}
created_doc_ids = []

# Pre-test cleanup of any previous test docs
try:
    existing_docs = client.get("/api/documents").json()
    for d in existing_docs:
        doc_id = d.get("document_id") or d.get("id")
        fname = d.get("filename", "")
        if doc_id and any(
            target in fname
            for target in [
                "test_chat_doc.pdf",
                "test_chat_doc.docx",
                "test_chat_invoice.jpg",
                "test_chat_receipt.png",
                "test_blank_image.jpg",
                "agentic_arch.jpg",
            ]
        ):
            client.delete(f"/api/documents/{doc_id}")
except Exception as e:
    print(f"Pre-cleanup notice: {e}")


# ============================================================================
# 1. PDF QUESTION ANSWERING
# ============================================================================
print("\n[TEST 1] Testing PDF Question Answering via /api/chat...")
doc_pdf = fitz.open()
page = doc_pdf.new_page()
page.insert_text(
    (50, 100),
    "DocuMind Vector Indexing: The indexing engine builds dense embeddings with all-MiniLM-L6-v2.\n"
    "Hybrid keyword matching uses BM25 algorithm with frequency normalization.",
    fontsize=12,
)
pdf_bytes = doc_pdf.tobytes()
doc_pdf.close()

pdf_upload = client.post(
    "/api/upload",
    files={"files": ("test_chat_doc.pdf", pdf_bytes, "application/pdf")},
)
assert pdf_upload.status_code == 200, f"PDF upload failed: {pdf_upload.text}"
pdf_id = pdf_upload.json()["document_id"]
created_doc_ids.append(pdf_id)

pdf_chat = client.post(
    "/api/chat",
    json={
        "document_id": pdf_id,
        "question": "What model is used for dense embeddings?",
    },
)
print(f" -> Status: {pdf_chat.status_code}")
assert pdf_chat.status_code == 200, f"PDF chat failed: {pdf_chat.text}"
pdf_chat_data = pdf_chat.json()
print(f" -> Answer: {pdf_chat_data.get('answer')}")
assert "MiniLM" in pdf_chat_data.get("answer", "") or "all-MiniLM-L6-v2" in pdf_chat_data.get("answer", "")
assert len(pdf_chat_data.get("sources", [])) > 0
print(f" -> Source: {pdf_chat_data['sources'][0]['title']}")
test_results["PDF Question Answering"] = "PASS"


# ============================================================================
# 2. DOCX QUESTION ANSWERING
# ============================================================================
print("\n[TEST 2] Testing DOCX Question Answering via /api/chat...")
docx = Document()
docx.add_paragraph("DocuMind AI Document Pipeline Architecture.")
docx.add_paragraph("The memory storage module caches conversation histories in JSON format.")
docx.add_paragraph("The maximum token limit for context is fourteen thousand characters.")
docx_buf = BytesIO()
docx.save(docx_buf)
docx_bytes = docx_buf.getvalue()

docx_upload = client.post(
    "/api/upload",
    files={"files": ("test_chat_doc.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
)
assert docx_upload.status_code == 200, f"DOCX upload failed: {docx_upload.text}"
docx_id = docx_upload.json()["document_id"]
created_doc_ids.append(docx_id)

docx_chat = client.post(
    "/api/chat",
    json={
        "document_id": docx_id,
        "question": "How does the memory storage module cache conversations?",
    },
)
print(f" -> Status: {docx_chat.status_code}")
assert docx_chat.status_code == 200, f"DOCX chat failed: {docx_chat.text}"
docx_chat_data = docx_chat.json()
print(f" -> Answer: {docx_chat_data.get('answer')}")
assert "JSON" in docx_chat_data.get("answer", "") or "json" in docx_chat_data.get("answer", "").lower()
assert len(docx_chat_data.get("sources", [])) > 0
print(f" -> Source: {docx_chat_data['sources'][0]['title']}")
test_results["DOCX Question Answering"] = "PASS"


# ============================================================================
# 3. JPG OCR QUESTION ANSWERING
# ============================================================================
print("\n[TEST 3] Testing JPG OCR Question Answering via /api/chat...")
jpg_img = Image.new("RGB", (900, 300), color="white")
draw = ImageDraw.Draw(jpg_img)
draw.text((20, 50), "Invoice Number: INV-88219. Total Amount: $4,250.00 USD. Due Date: 2026-11-15.", fill="black")
jpg_buf = BytesIO()
jpg_img.save(jpg_buf, format="JPEG")
jpg_bytes = jpg_buf.getvalue()

jpg_upload = client.post(
    "/api/upload",
    files={"files": ("test_chat_invoice.jpg", jpg_bytes, "image/jpeg")},
)
assert jpg_upload.status_code == 200, f"JPG upload failed: {jpg_upload.text}"
jpg_id = jpg_upload.json()["document_id"]
created_doc_ids.append(jpg_id)

jpg_chat = client.post(
    "/api/chat",
    json={
        "document_id": jpg_id,
        "question": "What is the invoice number and total amount?",
    },
)
print(f" -> Status: {jpg_chat.status_code}")
assert jpg_chat.status_code == 200, f"JPG chat failed: {jpg_chat.text}"
jpg_chat_data = jpg_chat.json()
print(f" -> Answer: {jpg_chat_data.get('answer')}")
assert "INV-88219" in jpg_chat_data.get("answer", "") or "4,250" in jpg_chat_data.get("answer", "") or "4250" in jpg_chat_data.get("answer", "")
assert len(jpg_chat_data.get("sources", [])) > 0
print(f" -> Source: {jpg_chat_data['sources'][0]['title']}")
test_results["JPG OCR Question Answering"] = "PASS"


# ============================================================================
# 4. PNG OCR QUESTION ANSWERING
# ============================================================================
print("\n[TEST 4] Testing PNG OCR Question Answering via /api/chat...")
png_img = Image.new("RGB", (900, 300), color="white")
draw = ImageDraw.Draw(png_img)
draw.text((20, 50), "Security Policy: Authentication tokens expire every 60 minutes. Encryption standard: AES-256.", fill="black")
png_buf = BytesIO()
png_img.save(png_buf, format="PNG")
png_bytes = png_buf.getvalue()

png_upload = client.post(
    "/api/upload",
    files={"files": ("test_chat_receipt.png", png_bytes, "image/png")},
)
assert png_upload.status_code == 200, f"PNG upload failed: {png_upload.text}"
png_id = png_upload.json()["document_id"]
created_doc_ids.append(png_id)

png_chat = client.post(
    "/api/chat",
    json={
        "document_id": png_id,
        "question": "What is the encryption standard and token expiration time?",
    },
)
print(f" -> Status: {png_chat.status_code}")
assert png_chat.status_code == 200, f"PNG chat failed: {png_chat.text}"
png_chat_data = png_chat.json()
print(f" -> Answer: {png_chat_data.get('answer')}")
assert "AES-256" in png_chat_data.get("answer", "") or "60 minutes" in png_chat_data.get("answer", "")
assert len(png_chat_data.get("sources", [])) > 0
print(f" -> Source: {png_chat_data['sources'][0]['title']}")
test_results["PNG OCR Question Answering"] = "PASS"


# ============================================================================
# 5. EMPTY OCR TEXT HANDLING
# ============================================================================
print("\n[TEST 5] Testing Empty OCR Text Handling via /api/chat...")
blank_img = Image.new("RGB", (800, 400), color="white")  # completely white canvas, no text
blank_buf = BytesIO()
blank_img.save(blank_buf, format="JPEG")
blank_bytes = blank_buf.getvalue()

blank_upload = client.post(
    "/api/upload",
    files={"files": ("test_blank_image.jpg", blank_bytes, "image/jpeg")},
)
assert blank_upload.status_code == 200, f"Blank image upload failed: {blank_upload.text}"
blank_id = blank_upload.json()["document_id"]
created_doc_ids.append(blank_id)

blank_chat = client.post(
    "/api/chat",
    json={
        "document_id": blank_id,
        "question": "What text is shown in this image?",
    },
)
print(f" -> Status: {blank_chat.status_code}")
assert blank_chat.status_code == 200, f"Expected 200 response for empty OCR image, got {blank_chat.status_code}: {blank_chat.text}"
blank_chat_data = blank_chat.json()
print(f" -> Answer: {blank_chat_data.get('answer')}")
assert "No readable text was detected in this image." in blank_chat_data.get("answer", ""), f"Unexpected answer: {blank_chat_data.get('answer')}"
test_results["Empty OCR Text"] = "PASS"


# ============================================================================
# 6. MISSING DOCUMENT ID HANDLING
# ============================================================================
print("\n[TEST 6] Testing Missing Document ID Handling...")
no_doc_chat = client.post(
    "/api/chat",
    json={
        "question": "What is the summary of this document?",
    },
)
print(f" -> Status: {no_doc_chat.status_code}")
print(f" -> Detail: {no_doc_chat.json().get('detail')}")
assert no_doc_chat.status_code == 400, f"Expected 400 for missing document ID, got {no_doc_chat.status_code}"
assert "Document ID is required" in no_doc_chat.json().get("detail", "")
test_results["Missing Document ID"] = "PASS"


# ============================================================================
# 7. MISSING QUESTION HANDLING
# ============================================================================
print("\n[TEST 7] Testing Missing Question Handling...")
no_q_chat = client.post(
    "/api/chat",
    json={
        "document_id": pdf_id,
        "question": "   ",
    },
)
print(f" -> Status: {no_q_chat.status_code}")
print(f" -> Detail: {no_q_chat.json().get('detail')}")
assert no_q_chat.status_code == 400, f"Expected 400 for missing question, got {no_q_chat.status_code}"
assert "Question is required" in no_q_chat.json().get("detail", "")
test_results["Missing Question"] = "PASS"


# ============================================================================
# 8. AGENTIC ARCHITECTURE USER FLOW
# ============================================================================
print("\n[TEST 8] Testing 'what is all agentic architecture?' on uploaded image...")
arch_img = Image.new("RGB", (900, 300), color="white")
draw = ImageDraw.Draw(arch_img)
draw.text(
    (20, 50),
    "All Agentic Architecture: An agentic architecture uses autonomous AI agents that coordinate to solve complex workflows.",
    fill="black",
)
arch_buf = BytesIO()
arch_img.save(arch_buf, format="JPEG")
arch_bytes = arch_buf.getvalue()

arch_upload = client.post(
    "/api/upload",
    files={"files": ("agentic_arch.jpg", arch_bytes, "image/jpeg")},
)
assert arch_upload.status_code == 200
arch_id = arch_upload.json()["document_id"]
created_doc_ids.append(arch_id)

arch_chat = client.post(
    "/api/chat",
    json={
        "document_id": arch_id,
        "question": "what is all agentic architecture?",
    },
)
print(f" -> Status: {arch_chat.status_code}")
assert arch_chat.status_code == 200, f"Agentic architecture query failed: {arch_chat.text}"
arch_data = arch_chat.json()
print(f" -> Answer: {arch_data.get('answer')}")
assert "autonomous" in arch_data.get("answer", "").lower() or "agent" in arch_data.get("answer", "").lower()
assert len(arch_data.get("sources", [])) > 0
test_results["Agentic Architecture Q&A"] = "PASS"


# ============================================================================
# CLEANUP
# ============================================================================
print("\n[CLEANUP] Deleting test documents...")
for doc_id in created_doc_ids:
    client.delete(f"/api/documents/{doc_id}")
print(" -> Cleanup complete.")

print("\n" + "=" * 65)
print("ALL TEST SUITE CHECKS COMPLETED SUCCESSFULLY!")
print("=" * 65)
for test_name, status in test_results.items():
    print(f"  [✔] {test_name}: {status}")
print("=" * 65)
